import base64
from dataclasses import dataclass
from pathlib import Path

import docx

_IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


class UnsupportedResumeFormat(Exception):
    pass


@dataclass
class ResumeContent:
    """What to send Claude for a given resume, plus how we got it."""

    extraction_method: str  # "text" | "vision"
    content_blocks: list[dict]  # Claude message content blocks to include after the prompt text
    # Pages the resume parser processed — always >= 1 (every resume is one parse). PDFs use the real
    # page count; images and text documents count as at least one page (text docs are estimated from
    # their length). Feeds the admin usage/cost dashboard's resume-parser metering.
    pages: int = 1


_CHARS_PER_PAGE = 2800  # rough single-page character budget for estimating text-document pages


def _pdf_page_count(path: Path) -> int:
    """Best-effort page count for metering — a broken/encrypted PDF still bills as 1 page."""
    try:
        from pypdf import PdfReader

        return max(1, len(PdfReader(str(path)).pages))
    except Exception:  # noqa: BLE001
        return 1


def _estimate_text_pages(text: str) -> int:
    """A text document has no rendered page count, so estimate one from its length (min 1)."""
    return max(1, round(len(text) / _CHARS_PER_PAGE))


def build_resume_content(stored_path: str) -> ResumeContent:
    path = Path(stored_path)
    if not path.exists():
        raise UnsupportedResumeFormat(f"Resume file not found on disk: {stored_path}")

    extension = path.suffix.lower()

    if extension == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode("utf-8")
        return ResumeContent(
            extraction_method="vision",
            content_blocks=[
                {
                    "type": "document",
                    "source": {"type": "base64", "media_type": "application/pdf", "data": data},
                }
            ],
            pages=_pdf_page_count(path),
        )

    if extension in _IMAGE_MEDIA_TYPES:
        data = base64.standard_b64encode(path.read_bytes()).decode("utf-8")
        return ResumeContent(
            extraction_method="vision",
            content_blocks=[
                {
                    "type": "image",
                    "source": {"type": "base64", "media_type": _IMAGE_MEDIA_TYPES[extension], "data": data},
                }
            ],
            pages=1,
        )

    if extension == ".docx":
        text = _extract_docx_text(path)
        if not text.strip():
            raise UnsupportedResumeFormat("The .docx file has no extractable text (it may be empty or image-only).")
        return ResumeContent(
            extraction_method="text",
            content_blocks=[{"type": "text", "text": f"Resume text:\n\n{text}"}],
            pages=_estimate_text_pages(text),
        )

    if extension == ".doc":
        raise UnsupportedResumeFormat(
            "Legacy .doc format is not supported — ask the candidate/recruiter to re-save as .docx or PDF."
        )

    raise UnsupportedResumeFormat(f"Unsupported file extension '{extension}'")


def _extract_docx_text(path: Path) -> str:
    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)
